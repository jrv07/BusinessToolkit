import logging
import os
import shutil
import share_objects as share
import platform
from pptx.util import Cm, Pt
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT as PP_ALIGN
from pptx.enum.text import MSO_VERTICAL_ANCHOR as MSO_ANCHOR


def setup_logger():
    # Log file
    # logFile = share.mainFileName.replace('.py','') + time.strftime('%Y%m%d_%H%M%S', time.localtime()) + '.log'
    log_file = 'analysis_toolkit.log'
    # Clear the content of logFile
    open(log_file, 'w').close()
    # Set parameters to write into logFile
    log_level = logging.DEBUG
    # logLevel = logging.INFO
    # logLevel = logging.ERROR
    logging.basicConfig(filename=log_file, filemode='a', level=log_level, datefmt='%a, %d %b %Y %H:%M:%S',
                        format='%(asctime)s %(name)-12s %(levelname)-8s %(message)s')
    # define a Handler which writes INFO messages or higher to the sys.stderr
    console = logging.StreamHandler()
    console.setLevel(logging.INFO)
    # set a format which is simpler for console use
    formatter = logging.Formatter('%(levelname)-8s: %(message)s')
    # tell the handler to use this format
    console.setFormatter(formatter)
    # add the handler to the root logger
    logging.getLogger('').addHandler(console)


def read_current_system():
    share.current_platform_system = platform.system()


def create_directories():
    #remove_main_folder(share.softwareDir)
    create_directory(share.softwareDir)
    create_directory(share.plotsDir)
    create_directory(share.curvesDir)
    create_directory(share.picturesDir)
    create_directory(share.moviesDir)
    create_directory(share.valuesDir)
    create_directory(share.animatorDir)
    create_directory(share.metapostDir)


def create_directory(path):
    if not os.path.exists(path):
        try:
            os.mkdir(path)
        except OSError:
            logging.error('Creation of directory "{}" failed'.format(path))
        else:
            logging.info('Directory "{}" successfully created'.format(path))


def remove_main_folder(path):
    if os.path.exists(path):
        shutil.rmtree(path)
        logging.info('Directory "{}" successfully removed'.format(path))


def set_table_cell(cell, text, font_size=Cm(12), fg_color=None, bg_color=None, alignment=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.font.size = font_size

    if fg_color:
        p.font.color.rgb = fg_color
        p.alignment = alignment
        cell.vertical_anchor = anchor

    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color


def set_table_cell_font_color(cell, color):
    p = cell.text_frame.paragraphs[0]
    p.font.color.rgb = color


def set_cell(cell, text):
    set_table_cell(cell, text, font_size=Pt(12), alignment=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def write_file(filename: str, content: str):
    with open(filename, "w") as f:
        f.write(content)
