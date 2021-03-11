from pptx.util import Cm
from typing import List
from pptx.table import Table, _Row
from module_pptx.PptxShapes import PptxShape
from module_pptx.pptx_table.PptxTableRow import PptxTableRow


class PptxTableBase(PptxShape):
    def __init__(self, slide, parameters, settings=None):
        super().__init__(slide, parameters, settings)
        self._rows: List[PptxTableRow] = []
        self._column_widths: List[Cm] = []
        self._col_count = 0
        self._row_count = 0

    def generate_table(self):
        self._shape_object = self._slide.shapes.add_table(self._row_count, self._col_count,
                                                          self._left, self._top, self._width, self._height)
        # save Powerpoint as xml, look for uuid in <a:tableStyleId>UUID</a:tableStyleId>
        # Todo: add as property
        # style_id = '{69012ECD-51FC-41F1-AA8D-1B2483CD663E}'   # blue horizontal borders
        # style_id = '{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}'   # without borders
        style_id = '{5940675A-B579-460E-94D1-54222C63F5DA}'     # with borders around each cell
        tbl = self._shape_object._element.graphic.graphicData.tbl
        tbl[0][-1].text = style_id
        self._shape_object.table.first_row = False
        self.generate_content(self._shape_object.table)

    @staticmethod
    def get_cm_list(float_list):
        cm_list = []
        for float_item in float_list:
            try:
                item_in_cm = Cm(float(float_item))
            except ValueError:
                return None
            cm_list.append(item_in_cm)
        return cm_list

    def get_size(self):
        width = Cm(0)
        for col_width in self._column_widths:
            width += col_width
        height = Cm(0)
        for row in self._rows:
            height = height + row.height
        self._width = width
        self._height = height

    def generate_content(self, pptx_table: Table):
        for col_nr, pptx_col in enumerate(pptx_table.columns):
            pptx_col.width = self._column_widths[col_nr]

        for row_index, row in enumerate(self._rows):
            pptx_row: _Row = pptx_table.rows[row_index]
            pptx_row.height = row.height
            for col_index, row_cell in enumerate(row._row_cells):
                cell = pptx_table.cell(row_index, col_index)
                row_cell.apply_styles(cell)

    def get_arguments(self):
        super().get_arguments()

    def generate(self):
        super().generate()
        self.get_size()
